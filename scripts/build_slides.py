"""Build Slides_Group09.pptx (16:9, 18 slides, ~20 minutes) from the notebook outputs.

Numbers come from outputs/results.json and outputs/model_comparison.csv; charts from figures/.
Each slide carries speaker notes with the talking points and the key numbers.
"""
import json
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Cm, Pt, Emu

ROOT = Path(__file__).resolve().parent.parent
R = json.loads((ROOT / "outputs" / "results.json").read_text())
CMP = pd.read_csv(ROOT / "outputs" / "model_comparison.csv").set_index("Model")
FIG = ROOT / "figures"
GROUP = "Group09"

NAVY = RGBColor(0x10, 0x42, 0x81)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
INK = RGBColor(0x0B, 0x0B, 0x0B)
GREY = RGBColor(0x52, 0x51, 0x4E)
LIGHT = RGBColor(0xEE, 0xF4, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
slide_no = [0]


def money(x, d=0):
    return f"${x:,.{d}f}"


def text(slide, x, y, w, h, content, size=16, bold=False, colour=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """Add a text box. `content` is a string or a list of (text, size, bold, colour) paragraphs."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(0.1); tf.margin_top = tf.margin_bottom = Cm(0.05)
    paras = content if isinstance(content, list) else [(content, size, bold, colour)]
    for i, item in enumerate(paras):
        t, s, b, c = (item + (None,) * 4)[:4] if isinstance(item, tuple) else (item, None, None, None)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        r = p.add_run(); r.text = t
        r.font.name = FONT; r.font.size = Pt(s or size); r.font.bold = bold if b is None else b
        r.font.color.rgb = c or colour
    return tb


def new_slide(title, kicker=None, notes=""):
    s = prs.slides.add_slide(BLANK)
    slide_no[0] += 1
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Cm(0.45), SH)
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    if kicker:
        text(s, Cm(1.5), Cm(0.8), Cm(30), Cm(0.9), kicker.upper(), size=12, bold=True, colour=BLUE)
    text(s, Cm(1.5), Cm(1.5), Cm(31), Cm(2.0), title, size=26, bold=True, colour=INK)
    text(s, Cm(1.5), SH - Cm(1.1), Cm(20), Cm(0.8), f"ACC2125 {GROUP} · Sydney short-term rentals", size=10, colour=GREY)
    text(s, SW - Cm(3.5), SH - Cm(1.1), Cm(2.5), Cm(0.8), str(slide_no[0]), size=10, colour=GREY, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


def picture(slide, name, x, y, max_w, max_h):
    """Place a figure scaled to fit within max_w x max_h, centred in that box."""
    w_px, h_px = Image.open(FIG / f"{name}.png").size
    scale = min(max_w / w_px, max_h / h_px)
    w, h = int(w_px * scale), int(h_px * scale)
    slide.shapes.add_picture(str(FIG / f"{name}.png"), x + (max_w - w) // 2, y + (max_h - h) // 2, w, h)


def takeaways(slide, items, x, y, w, h, size=16):
    """Bulleted takeaways: list of (bold lead, rest)."""
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    for i, (lead, rest) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        r = p.add_run(); r.text = lead + " "; r.font.bold = True; r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = NAVY
        r = p.add_run(); r.text = rest; r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = INK
    return tb


def stat(slide, x, y, value, label, w=Cm(7), colour=NAVY):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Cm(3.2))
    box.fill.solid(); box.fill.fore_color.rgb = LIGHT; box.line.fill.background()
    text(slide, x + Cm(0.4), y + Cm(0.25), w - Cm(0.8), Cm(1.6), value, size=30, bold=True, colour=colour)
    text(slide, x + Cm(0.4), y + Cm(1.85), w - Cm(0.8), Cm(1.3), label, size=12, colour=GREY)


def table(slide, df, x, y, w, col_w_cm, size=13):
    rows, cols = df.shape[0] + 1, df.shape[1]
    shp = slide.shapes.add_table(rows, cols, x, y, w, Cm(0.9) * rows)
    t = shp.table
    for j, cw in enumerate(col_w_cm):
        t.columns[j].width = Cm(cw)
    for j, c in enumerate(df.columns):
        cell = t.cell(0, j); cell.text = str(c)
        cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            cell = t.cell(i, j); cell.text = str(v)
            cell.fill.solid(); cell.fill.fore_color.rgb = LIGHT if i % 2 == 0 else WHITE
            for p in cell.text_frame.paragraphs:
                if j > 0 and str(v)[:1] in "0123456789$-+.":
                    p.alignment = PP_ALIGN.RIGHT
                for r in p.runs:
                    r.font.size = Pt(size); r.font.color.rgb = INK; r.font.name = FONT
    return t


q1 = pd.DataFrame(R["q1"]); pop, oth = q1["Popular (top 25%)"], q1["Other listings"]
m1, m3 = R["m1"], R["m3"]
pr = pd.DataFrame(R["m2_profiles"])
city = pr[pr.segment == "City-core high-turnover"].iloc[0]; ring = pr[pr.segment == "Inner-ring low-utilisation"].iloc[0]
ht = pd.DataFrame(R["a3_host_tier"]).set_index("host_tier"); hs = R["a3_host_subscores"]; tt = R["a3_host_tests"]
S, C = ht.loc["Single (1)"], ht.loc["Commercial (10+)"]
bands = R["q3_bands"]; mn = R["q1_occ_by_minnights"]; rb = R["q1_occ_by_rating"]; sh = R["q1_occ_superhost"]

# ---------------------------------------------------------------- 1. Title
s = prs.slides.add_slide(BLANK); slide_no[0] += 1
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH); bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
text(s, Cm(2.5), Cm(3.0), Cm(29), Cm(1), "ACC2125 DATA ANALYTICS AND MACHINE LEARNING · GROUP PROJECT", size=14, bold=True, colour=RGBColor(0x9E, 0xC5, 0xF4))
text(s, Cm(2.5), Cm(4.5), Cm(29), Cm(3), "Does Rail Access Pay?", size=48, bold=True, colour=WHITE)
text(s, Cm(2.5), Cm(7.3), Cm(29), Cm(2), "Machine learning on Sydney's short-term rental market", size=24, colour=WHITE)
text(s, Cm(2.5), Cm(11.0), Cm(29), Cm(4.5),
     [(f"{GROUP}", 18, True, WHITE),
      ("[Member 1] · [Member 2] · [Member 3] · [Member 4]", 16, False, WHITE),
      (f"Inside Airbnb Sydney snapshot (June 2026) + Transport for NSW station entrances", 14, False, RGBColor(0x9E, 0xC5, 0xF4))])
s.notes_slide.notes_text_frame.text = (
    "Introduce the group and the question: Sydney has a dense rail network and a huge Airbnb market. Investors often assume "
    "listings near stations earn more. We test that with Inside Airbnb data joined to every Transport for NSW station entrance. "
    "Spoiler: rail access buys bookings, not higher prices. (~30 s)")

# ---------------------------------------------------------------- 2. Agenda / question
s = new_slide("Our question: what drives price and demand in Sydney?", "Overview",
              "Frame the talk. Three parts: what the data shows (guided Q1-Q3 + our two own analyses), what the models show, "
              "and what hosts and investors should do. The external rail data is used in every part. Hand over between sections. (~45 s)")
takeaways(s, [("Part A — Exploration:", "who gets booked (Q1), does price buy ratings (Q2), how location matters (Q3)."),
              ("Unguided analyses:", "(1) is there a rail-access price premium? (2) do commercial hosts deliver worse stays?"),
              ("Part B — Machine learning:", "Ridge regression, K-Means segmentation, XGBoost + SHAP (self-taught)."),
              ("Part C — Recommendations:", "for investors, individual hosts and commercial operators."),
              ("Thread throughout:", "Transport for NSW station-entrance data, joined to every listing.")],
          Cm(1.5), Cm(4.2), Cm(30), Cm(12), size=20)

# ---------------------------------------------------------------- 3. Data sample
s = new_slide("Data sample: Sydney, June 2026 snapshot", "A1 · Data selection",
              f"Inside Airbnb Sydney, scraped 17-29 June 2026, CC BY 4.0. {R['raw_listings']:,} listings; after cleaning {R['clean_listings']:,}. "
              "Main cleaning steps: drop listings with no price, drop mid/long-term rentals with minimum stay over 30 nights, trim the 1st and 99th "
              f"price percentiles ({money(R['price_trim'][0])}-{money(R['price_trim'][1])}), impute bedrooms/beds from similar listings, keep ratings "
              "missing rather than invent them. Every step is logged in the Data Decisions Log in the report. "
              "Why Sydney: big, mature market with very different sub-markets and a dense rail network. (~1 min)")
stat(s, Cm(1.5), Cm(4.3), f"{R['raw_listings']:,}", "raw listings (Inside Airbnb)")
stat(s, Cm(9.0), Cm(4.3), f"{R['clean_listings']:,}", "short-term listings analysed")
stat(s, Cm(16.5), Cm(4.3), f"{R['raw_calendar_rows']/1e6:.1f} M", "calendar rows (next 12 months)")
stat(s, Cm(24.0), Cm(4.3), "38", "local government areas (LGAs)")
takeaways(s, [("Why Sydney:", "a large market with distinct sub-markets (CBD, beaches, west) and dense rail, so we can test transport effects."),
              ("Cleaning:", f"removed {R['raw_listings'] - R['clean_listings']:,} listings: no price, 30+ night minimum stays, extreme prices (1st/99th pct)."),
              ("Demand measure:", "Inside Airbnb's estimated nights booked (12 months), checked against the 90-day calendar."),
              ("Documented:", "11-step Data Decisions Log in the report and notebook.")],
          Cm(1.5), Cm(8.5), Cm(30.5), Cm(8.5), size=17)

# ---------------------------------------------------------------- 4. External data
s = new_slide("External data: every rail station entrance in NSW", "A1 · External dataset",
              f"Transport for NSW 'Train station entrance locations' (2020, v4), Open Data Hub, CC BY 4.0, accessed 6 Oct 2026. "
              f"{R['raw_station_rows']:,} entrances; we keep {R['entrances_sydney']} in Greater Sydney ({R['stations_sydney']} stations incl. light rail and Metro). "
              "Entrances rather than station centroids, because the walk from the door matters. We use a haversine BallTree to find the nearest entrance "
              f"for each listing and count stations within 1 km. Median listing is {R['median_dist_station_m']} m from an entrance; "
              f"{R['share_within_500m']}% are within 500 m. These features feed Q3, both unguided analyses and all three models. (~1 min)")
picture(s, "q3_lga_maps", Cm(1.2), Cm(3.8), Cm(17.5), Cm(13.8))
takeaways(s, [("Source:", "Transport for NSW Open Data Hub, CC BY 4.0 (2020 network)."),
              ("Coverage:", f"{R['entrances_sydney']} entrances at {R['stations_sydney']} Sydney stations (black dots)."),
              ("Integration:", "haversine nearest-neighbour join to every listing."),
              ("Features:", "distance to nearest entrance, stations within 1 km, distance band."),
              ("Result:", f"median {R['median_dist_station_m']} m; {R['share_within_500m']}% of listings within 500 m.")],
          Cm(19.2), Cm(4.2), Cm(13.5), Cm(13), size=16)

# ---------------------------------------------------------------- 5. Q1
s = new_slide("Q1 · Popular listings win on operation, not price", "A2 · Guided analysis",
              f"Popular = top 25% by estimated nights booked (>= {R['popular_threshold_nights']:.0f} nights). They are cheaper "
              f"({money(pop['Median price (AUD)'])} vs {money(oth['Median price (AUD)'])}). {pop['Superhost %']:.0f}% are superhosts vs {oth['Superhost %']:.0f}%. "
              f"Occupancy drops from {mn['1']:.0f} nights at 1-night minimum to {mn['8-30']:.0f} at 8-30. Ratings above 4.5 roughly double occupancy. "
              "Caveat: occupancy is estimated from reviews, so it is a proxy. (~1 min)")
picture(s, "q1_popularity_drivers", Cm(1.2), Cm(3.9), Cm(31.5), Cm(9.5))
takeaways(s, [("Cheaper, not dearer:", f"popular listings median {money(pop['Median price (AUD)'])} vs {money(oth['Median price (AUD)'])}."),
              ("Superhosts:", f"{pop['Superhost %']:.0f}% of popular listings vs {oth['Superhost %']:.0f}% of others."),
              ("Flexibility:", f"{mn['1']:.0f} nights booked at 1-night minimum vs {mn['8-30']:.0f} at 8-30 nights.")],
          Cm(1.5), Cm(13.6), Cm(31), Cm(4), size=16)

# ---------------------------------------------------------------- 6. Q2
dec = R["q2_deciles"]; sp = R["q2_spearman_price"]
s = new_slide("Q2 · Higher prices buy only slightly better ratings", "A2 · Guided analysis",
              f"{R['q2_n_rated']:,} listings with 5+ reviews. Spearman rho = {R['q2_rho_overall']:.2f}: positive but weak. "
              f"Non-linear: the cheapest decile averages {dec['rating_mean'][0]:.2f} stars; by the second decile it is {dec['rating_mean'][1]:.2f}, then flat. "
              f"Location sub-score tracks price most (rho {sp['Location']:.2f}); value barely (rho {sp['Value']:.2f}). "
              f"Log price explains only {100*R['q2_linear_r2']:.1f}% of rating variation. (~1 min)")
picture(s, "q2_price_vs_rating", Cm(1.2), Cm(3.9), Cm(31.5), Cm(10))
takeaways(s, [("Weak link:", f"ρ = {R['q2_rho_overall']:.2f}; price explains {100*R['q2_linear_r2']:.0f}% of rating variation."),
              ("Floor effect:", f"{dec['share_below_45'][0]:.0f}% of the cheapest decile rated < 4.5 vs {dec['share_below_45'][9]:.0f}% of the dearest."),
              ("Value ≠ price:", f"value score barely moves with price (ρ = {sp['Value']:.2f}).")],
          Cm(1.5), Cm(14.1), Cm(31), Cm(3.5), size=16)

# ---------------------------------------------------------------- 7. Q3 LGA
tp = R["q3_lga_top_price"]; to = R["q3_lga_top_occ"]
s = new_slide("Q3 · Price and demand follow different maps", "A2 · Guided analysis",
              f"Price peaks on the coast and North Shore: Pittwater {money(tp['Pittwater']['median_price'])}, Manly {money(tp['Manly']['median_price'])}. "
              f"Demand peaks in the City of Sydney ({to['Sydney']['median_occupancy']:.0f} nights). Pittwater books only {tp['Pittwater']['median_occupancy']:.0f}. "
              f"Across LGAs price and occupancy are nearly unrelated (rho {R['q3_lga_price_occ_rho']}). Leisure areas earn on rate, the city on volume. (~1 min)")
picture(s, "q3_lga_maps", Cm(1.2), Cm(3.8), Cm(20), Cm(13.8))
takeaways(s, [("Rate markets:", f"Pittwater {money(tp['Pittwater']['median_price'])}, Manly {money(tp['Manly']['median_price'])}, Woollahra {money(tp['Woollahra']['median_price'])}."),
              ("Volume market:", f"City of Sydney: {to['Sydney']['median_occupancy']:.0f} nights at {money(to['Sydney']['median_price'])}."),
              ("Concentration:", f"{R['q3_sydney_lga_share']}% of listings in one LGA."),
              ("Unrelated:", f"LGA price vs occupancy ρ = {R['q3_lga_price_occ_rho']}.")],
          Cm(21.5), Cm(4.5), Cm(11.5), Cm(13), size=16)

# ---------------------------------------------------------------- 8. Q3 rail
s = new_slide("Q3 · Near rail: more bookings, lower prices", "A2 · Guided analysis (external data)",
              f"Using the TfNSW data: within 500 m listings book {bands['median_occupancy'][0]:.0f} median nights vs {bands['median_occupancy'][3]:.0f} beyond 2 km, "
              f"but are cheaper ({money(bands['median_price'][0])} vs {money(bands['median_price'][3])}). The far band contains beach homes. "
              f"Rail distance and CBD distance are correlated (rho {R['q3_rho_cbd_station']}), so we need controls: that is Analysis 1. (~45 s)")
picture(s, "q3_transit_bands", Cm(1.2), Cm(3.9), Cm(31.5), Cm(9.8))
takeaways(s, [("Demand:", f"{bands['median_occupancy'][0]:.0f} vs {bands['median_occupancy'][3]:.0f} median nights (<500 m vs >2 km)."),
              ("Price:", f"{money(bands['median_price'][0])} vs {money(bands['median_price'][3])}: the far band includes beach houses."),
              ("Confounded:", f"rail and CBD distance correlate (ρ = {R['q3_rho_cbd_station']}), so we need controls.")],
          Cm(1.5), Cm(14.0), Cm(31), Cm(3.5), size=16)

# ---------------------------------------------------------------- 9. A3-1
s = new_slide("Analysis 1 · The 'transit premium' disappears under controls", "A3 · Unguided analysis",
              f"Log-price OLS with robust errors. Raw: {R['a3_raw_premium_500']:.0f}% within 500 m. Add listing structure: about -13%. "
              f"Add CBD distance + LGA: {R['a3_adj_premium_500']:+.1f}% (95% CI {R['a3_adj_premium_500_ci'][0]:+.1f} to {R['a3_adj_premium_500_ci'][1]:+.1f}), "
              f"about {money(R['a3_adj_premium_500_dollars'])}/night, not significant. But occupancy: +{R['a3_occ_500_nights']:.1f} nights/year (p = {R['a3_occ_500_p']:.2f}). "
              "Message: rail access is an occupancy story, not a pricing story. (~1.5 min)")
picture(s, "a3_transit_premium", Cm(1.2), Cm(3.9), Cm(20.5), Cm(11))
stat(s, Cm(22.5), Cm(4.3), f"{R['a3_raw_premium_500']:+.0f}%", "raw price gap, < 500 m vs > 2 km", w=Cm(10), colour=ORANGE)
stat(s, Cm(22.5), Cm(7.9), f"{R['a3_adj_premium_500']:+.1f}%", f"after controls (CI {R['a3_adj_premium_500_ci'][0]:+.1f} to {R['a3_adj_premium_500_ci'][1]:+.1f}), not significant", w=Cm(10))
stat(s, Cm(22.5), Cm(11.5), f"+{R['a3_occ_500_nights']:.1f}", f"extra nights booked a year (p = {R['a3_occ_500_p']:.2f})", w=Cm(10), colour=BLUE)
text(s, Cm(1.5), Cm(15.5), Cm(31), Cm(1.5), "Controls: room type, guests, bedrooms, bathrooms, amenities, distance to CBD, LGA fixed effects (HC3 robust errors).", size=13, colour=GREY)

# ---------------------------------------------------------------- 10. A3-2
s = new_slide("Analysis 2 · Commercial hosts: best locations, weaker stays", "A3 · Unguided analysis",
              f"{C['hosts']:.0f} commercial hosts (10+ listings) run {C['listing_share_pct']:.0f}% of listings. Mean rating {C['mean_rating_5plus']:.2f} vs {S['mean_rating_5plus']:.2f} for single hosts; "
              f"cleanliness has the biggest gap. After controls still {abs(tt['rating_gap_adj']):.2f} stars lower. They are much closer to rail "
              f"({C['median_dist_station_m']:,.0f} m vs {S['median_dist_station_m']:,.0f} m) yet book fewer nights ({C['median_occupancy']:.0f} vs {S['median_occupancy']:.0f}). (~1.5 min)")
picture(s, "a3_host_tiers", Cm(1.2), Cm(3.9), Cm(31.5), Cm(9.8))
takeaways(s, [("Scale:", f"{C['hosts']:.0f} commercial hosts run {C['listing_share_pct']:.0f}% of listings."),
              ("Ratings:", f"{C['mean_rating_5plus']:.2f} vs {S['mean_rating_5plus']:.2f}; still {abs(tt['rating_gap_adj']):.2f} lower after controls; {C['below_4_5_pct']:.0f}% below 4.5."),
              ("Paradox:", f"closest to rail ({C['median_dist_station_m']:,.0f} m vs {S['median_dist_station_m']:,.0f} m) but fewest nights ({C['median_occupancy']:.0f} vs {S['median_occupancy']:.0f}).")],
          Cm(1.5), Cm(14.0), Cm(31), Cm(3.5), size=16)

# ---------------------------------------------------------------- 11. ML setup
s = new_slide("Machine learning set-up", "B · Machine learning",
              f"Same cleaned data for all models. Supervised target: log nightly price (right-skewed). 80/20 split, seed 42: {R['ml_train_n']:,} train, "
              f"{R['ml_test_n']:,} test. Tuning by cross-validation on training data only. Rail features from the external data are inputs in every model. "
              f"We also re-run each price model without the rail features (ablation) to measure what the external data adds. (~45 s)")
rows = [("1 · Ridge regression", "Supervised (in class)", "Predict log price; interpretable baseline", "α by 5-fold CV"),
        ("2 · K-Means", "Unsupervised (in class)", "Segment market into investment zones", "k by elbow + silhouette"),
        ("3 · XGBoost + SHAP", "Self-taught", "Best accuracy; explain rail effect", "30-config random search, 3-fold CV")]
table(s, pd.DataFrame(rows, columns=["Model", "Type", "Objective", "Tuning"]), Cm(1.5), Cm(4.3), Cm(30.5), [6.5, 6, 10, 8], size=15)
takeaways(s, [("Inputs:", "size, room type, property type, amenities, reviews, host scale, coordinates, LGA, distance to CBD, distance to station, stations within 1 km."),
              ("Evaluation:", f"held-out test set ({R['ml_test_n']:,} listings): R², RMSE, MAE, MAPE against a median-price baseline."),
              ("Ablation:", "each price model refitted without rail features to measure what the external data adds.")],
          Cm(1.5), Cm(9.8), Cm(30.5), Cm(7), size=17)

# ---------------------------------------------------------------- 12. Ridge
s = new_slide(f"Model 1 · Ridge regression explains {100*m1['R2 (log price)']:.0f}% of log-price variation", "B · Supervised",
              f"Ridge = linear regression + L2 penalty, good with correlated inputs. Best alpha {R['m1_alpha']:.0f}. Test R² {m1['R2 (log price)']:.3f}, "
              f"MAE {money(m1['MAE (AUD)'])}, MAPE {m1['MAPE %']:.1f}% (baseline MAPE {CMP.loc['Baseline (median)','MAPE %']:.0f}%). No overfitting (train {R['m1_train_r2']:.3f}). "
              f"Room type and size dominate. Rail coefficients tiny; dropping them: R² {m1['R2 (log price)']:.4f} -> {R['m1_noT']['R2 (log price)']:.4f}. "
              "Issues: rare room types merged; LGA dummies collinear with distance, so don't over-read location coefficients. (~1 min)")
picture(s, "m1_ridge", Cm(1.2), Cm(3.9), Cm(31.5), Cm(10.3))
takeaways(s, [("Performance:", f"R² {m1['R2 (log price)']:.2f}, MAE {money(m1['MAE (AUD)'])}, MAPE {m1['MAPE %']:.0f}% (baseline {CMP.loc['Baseline (median)','MAPE %']:.0f}%)."),
              ("Drivers:", "room type and size; rail coefficients near zero."),
              ("Issue fixed:", "rare room types merged; collinear location terms are interpreted with care.")],
          Cm(1.5), Cm(14.3), Cm(31), Cm(3.3), size=15)

# ---------------------------------------------------------------- 13. K-Means
s = new_slide("Model 2 · K-Means finds four investment zones", "B · Unsupervised",
              f"Inputs: coordinates, log price, log station distance, CBD distance, occupancy, all standardised. Silhouette peaks at k=3 "
              f"({R['m2_k_table']['silhouette'][1]:.2f}) but k=4 ({R['m2_silhouette']:.2f}) splits the inner city into two very different groups, "
              f"so we chose it. Key insight: the city-core and inner-ring segments share locations and rail access, but occupancy is "
              f"{city.median_occupancy:.0f} vs {ring.median_occupancy:.0f} nights and revenue {money(city.median_revenue)} vs {money(ring.median_revenue)}. (~1.5 min)")
picture(s, "m2_cluster_maps", Cm(1.2), Cm(3.8), Cm(31.5), Cm(7.5))
prow = [(r_.segment, f"{r_.listings:,.0f}", money(r_.median_price), f"{r_.median_occupancy:.0f}", money(r_.median_revenue), f"{r_.median_dist_station_m:,.0f} m")
        for r_ in pr.itertuples()]
table(s, pd.DataFrame(prow, columns=["Segment", "Listings", "Price", "Nights", "Revenue (12 m)", "To station"]),
      Cm(1.5), Cm(11.6), Cm(30.5), [10, 3.5, 3.5, 3.5, 5, 5], size=13)

# ---------------------------------------------------------------- 14. XGBoost + SHAP
sb = R["m3_shap_dist_band_pct"]
s = new_slide("Model 3 · XGBoost + SHAP: rail ranks 19th of 21 inputs", "B · Self-taught model",
              "XGBoost: gradient-boosted trees, each tree fits the previous errors; regularised. SHAP: Shapley-value attribution of each prediction. "
              f"Test R² {m3['R2 (log price)']:.3f}, MAPE {m3['MAPE %']:.1f}%. Bedrooms, guests, room type, longitude dominate. "
              f"Rail features: {R['m3_shap_rail_share_pct']}% of total SHAP; station distance effect within about ±1% below 4 km, "
              f"+{sb['>5 km']:.1f}% beyond 5 km (Northern Beaches proxy). Issues: train/test gap (0.91 vs 0.81) controlled by subsampling and regularisation; "
              "SHAP for one-hot columns summed back to each variable. (~1.5 min)")
picture(s, "m3_shap", Cm(1.2), Cm(3.9), Cm(31.5), Cm(10.6))
takeaways(s, [("Top drivers:", "bedrooms, guests, room type, longitude (coast vs inland)."),
              ("Rail:", f"{R['m3_shap_rail_share_pct']}% of total SHAP; ±1% effect within 4 km of a station."),
              ("Issue fixed:", "train/test gap controlled with subsampling + regularisation; one-hot SHAP regrouped.")],
          Cm(1.5), Cm(14.6), Cm(31), Cm(3), size=15)

# ---------------------------------------------------------------- 15. Comparison
s = new_slide("Model comparison: XGBoost wins; rail adds almost nothing", "B · Evaluation",
              f"XGBoost beats Ridge: R² {m3['R2 (log price)']:.3f} vs {m1['R2 (log price)']:.3f}; MAE {money(m3['MAE (AUD)'])} vs {money(m1['MAE (AUD)'])}. "
              "Ablation: removing rail features changes R² in the third decimal place for both. Three methods agree: regression with controls, "
              "Ridge, XGBoost/SHAP. Both models shrink to the middle: over-predict cheap, under-predict luxury. (~1 min)")
picture(s, "ml_comparison", Cm(1.2), Cm(3.9), Cm(31.5), Cm(8.6))
rows = [(n, f"{CMP.loc[n,'R2 (log price)']:.3f}", money(CMP.loc[n, 'MAE (AUD)']), f"{CMP.loc[n,'MAPE %']:.1f}%")
        for n in ["Model 1: Ridge", "Ridge without rail features", "Model 3: XGBoost", "XGBoost without rail features"]]
table(s, pd.DataFrame(rows, columns=["Model (test set)", "R²", "MAE", "MAPE"]), Cm(1.5), Cm(12.7), Cm(21), [10, 3.5, 3.5, 4], size=13)
text(s, Cm(23.5), Cm(12.8), Cm(9), Cm(4.2), [("Ablation verdict", 15, True, NAVY),
                                            ("Dropping rail features moves R² by ≤ 0.002 in both models.", 14, False, INK)])

# ---------------------------------------------------------------- 16. Recommendations
s = new_slide("Three recommendations", "C · Business recommendations",
              "1. Investors: no price premium near stations; value transit-adjacent units on occupancy (city-core segment). "
              "2. Hosts: cut minimum stays and work towards superhost criteria before cutting price. "
              f"3. Commercial operators: fix cleanliness and value; {C['below_4_5_pct']:.0f}% of their listings are below 4.5, where occupancy collapses. (~2 min)")
cards = [("1 · Investors", "Pay for occupancy, not for rate",
          f"No reliable price premium within 500 m of rail ({R['a3_adj_premium_500']:+.1f}%), but +{R['a3_occ_500_nights']:.0f} nights a year. "
          f"Underwrite transit-adjacent inner-city units on volume: city-core segment earns {money(city.median_revenue)} median revenue."),
         ("2 · Individual hosts", "Cut friction before cutting price",
          f"1-night minimums book {mn['1']:.0f} nights vs {mn['4-7']:.0f} at 4–7. Superhosts book ~{sh['1']/max(sh['0'],1):.0f}× more. "
          f"Price barely moves ratings (ρ = {R['q2_rho_overall']:.2f})."),
         ("3 · Commercial operators", "Close the guest-experience gap",
          f"{abs(tt['rating_gap_adj']):.2f} stars behind single hosts after controls; cleanliness gap largest. "
          f"{C['below_4_5_pct']:.0f}% of listings below 4.5, where occupancy falls to {rb['<4.5']:.0f} nights.")]
for i, (head, sub, body) in enumerate(cards):
    x = Cm(1.5 + i * 10.4)
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Cm(4.3), Cm(9.8), Cm(12.5))
    box.fill.solid(); box.fill.fore_color.rgb = LIGHT; box.line.fill.background()
    text(s, x + Cm(0.5), Cm(4.7), Cm(8.8), Cm(1.2), head, size=16, bold=True, colour=BLUE)
    text(s, x + Cm(0.5), Cm(5.9), Cm(8.8), Cm(2.4), sub, size=22, bold=True, colour=INK)
    text(s, x + Cm(0.5), Cm(8.8), Cm(8.8), Cm(7.5), body, size=16, colour=INK)

# ---------------------------------------------------------------- 17. Limitations
s = new_slide("Limitations and next steps", "C · Summary",
              "Be upfront: occupancy and revenue are Inside Airbnb estimates from reviews; prices are asking prices from one winter snapshot; "
              "the station file is the 2020 network, so new Metro stations are missing; results are associations. Next steps: multiple snapshots for seasonality, "
              "spatial cross-validation, the 2024 Metro network, and review text. (~1 min)")
takeaways(s, [("Demand is estimated:", "Inside Airbnb occupancy and revenue are modelled from reviews; the calendar mixes bookings with host blocks."),
              ("One snapshot:", "June (winter) asking prices; no seasonality or realised prices."),
              ("2020 rail network:", "newer Sydney Metro stations are not in the external dataset."),
              ("Associations, not causes:", "regression controls and SHAP explain patterns, not causal effects."),
              ("Next steps:", "multiple snapshots, spatial cross-validation, updated network data, NLP on review text.")],
          Cm(1.5), Cm(4.3), Cm(30.5), Cm(12.5), size=19)

# ---------------------------------------------------------------- 18. Close
s = new_slide("Rail access buys bookings, not higher prices", "Thank you · Questions",
              "Close with the one-line takeaway and invite questions. (~30 s)")
takeaways(s, [("Price:", "set by what a listing is and where it is; rail adds ≈ 0 after controls."),
              ("Demand:", "set by how it is run: flexibility, superhost standards, cleanliness, plus a modest rail boost."),
              ("Models:", f"XGBoost R² {m3['R2 (log price)']:.2f} vs Ridge {m1['R2 (log price)']:.2f}; K-Means separates high- and low-utilisation city stock.")],
          Cm(1.5), Cm(4.3), Cm(30.5), Cm(7.5), size=20)
text(s, Cm(1.5), Cm(12.6), Cm(30.5), Cm(4.2),
     [("Data: Inside Airbnb, Sydney (scraped 17–29 June 2026), CC BY 4.0 · Transport for NSW, Train station entrance locations (2020), CC BY 4.0", 12, False, GREY),
      ("Methods: Chen & Guestrin (2016) XGBoost · Lundberg & Lee (2017) SHAP · Hoerl & Kennard (1970) Ridge · Rousseeuw (1987) silhouette", 12, False, GREY)])

out = ROOT / f"Slides_{GROUP}.pptx"
prs.save(out)
print("Wrote", out, "| slides:", slide_no[0])
